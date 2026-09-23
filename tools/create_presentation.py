from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'Education-Center-MIS-Guide-Pashto.pptx'
ASSETS = ROOT / 'presentation_assets'
ASSETS.mkdir(exist_ok=True)

NAVY = RGBColor(9, 43, 70)
TEAL = RGBColor(11, 136, 126)
MINT = RGBColor(225, 248, 244)
INK = RGBColor(23, 51, 77)
MUTED = RGBColor(105, 126, 145)
AMBER = RGBColor(202, 137, 31)
WHITE = RGBColor(255, 255, 255)


def make_visual(name, kind):
    path = ASSETS / name
    image = Image.new('RGB', (1200, 700), '#eef7f7')
    draw = ImageDraw.Draw(image)
    if kind == 'dashboard':
        draw.rounded_rectangle((50, 45, 1150, 655), radius=30, fill='#ffffff', outline='#cfe1e8', width=4)
        draw.rectangle((50, 45, 290, 655), fill='#092b46')
        for y in (125, 195, 265, 335):
            draw.rounded_rectangle((85, y, 255, y + 32), 10, fill='#2cb6aa')
        for x, color in ((340, '#dff7f3'), (590, '#e7efff'), (840, '#fff2d9')):
            draw.rounded_rectangle((x, 95, x + 200, 205), 18, fill=color)
            draw.rectangle((x + 24, 125, x + 130, 142), fill='#0b887e')
            draw.rectangle((x + 24, 160, x + 85, 177), fill='#55718a')
        draw.rounded_rectangle((340, 255, 1110, 600), 18, fill='#f8fafc', outline='#d6e2ea', width=3)
        for y in range(300, 555, 48):
            draw.line((370, y, 1080, y), fill='#dce6ed', width=3)
        for x in (395, 620, 820):
            draw.rounded_rectangle((x, 325, x + 125, 350), 8, fill='#51cbbf')
    elif kind == 'install':
        draw.rounded_rectangle((230, 90, 970, 570), radius=35, fill='#092b46')
        draw.rounded_rectangle((285, 135, 915, 505), radius=15, fill='#ffffff')
        draw.rounded_rectangle((350, 185, 600, 255), 15, fill='#e0f8f4')
        draw.rounded_rectangle((350, 295, 825, 345), 12, fill='#e7efff')
        draw.rounded_rectangle((350, 380, 735, 430), 12, fill='#fff2d9')
        draw.polygon([(740, 185), (825, 185), (825, 135), (925, 235), (825, 335), (825, 285), (740, 285)], fill='#0b887e')
    elif kind == 'workflow':
        nodes = [(85, 280, '#0b887e'), (325, 280, '#2584cf'), (565, 280, '#c18115'), (805, 280, '#7f60b9')]
        for x, y, color in nodes:
            draw.ellipse((x, y, x + 150, y + 150), fill=color)
        for x in (235, 475, 715):
            draw.polygon([(x, 340), (x + 60, 315), (x + 60, 330), (x + 115, 330), (x + 115, 350), (x + 60, 350), (x + 60, 365)], fill='#55718a')
    else:
        draw.rounded_rectangle((260, 110, 940, 590), radius=30, fill='#ffffff', outline='#b6d8da', width=5)
        draw.rectangle((330, 190, 870, 240), fill='#0b887e')
        draw.rectangle((330, 285, 760, 330), fill='#dff7f3')
        draw.rectangle((330, 370, 830, 415), fill='#e7efff')
        draw.ellipse((730, 430, 860, 560), fill='#f4c96f')
    image.save(path)
    return path


def text_box(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.RIGHT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = 'Tahoma'
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = Inches(.08)
    return box


def add_header(slide, title, subtitle=''):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(245, 248, 252)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.22))
    shape.fill.solid(); shape.fill.fore_color.rgb = TEAL; shape.line.fill.background()
    text_box(slide, title, .6, .38, 12.1, .55, 25, NAVY, True)
    if subtitle:
        text_box(slide, subtitle, .6, .94, 12.1, .35, 12, MUTED)


def bullet_slide(prs, title, subtitle, bullets, visual=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, title, subtitle)
    top = 1.48
    width = 7.1 if visual else 11.8
    for index, (heading, body) in enumerate(bullets, start=1):
        y = top + (index - 1) * 1.13
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.9 - width + .1), Inches(y), Inches(.43), Inches(.43))
        badge.fill.solid(); badge.fill.fore_color.rgb = TEAL; badge.line.fill.background()
        text_box(slide, str(index), 11.9 - width + .1, y + .03, .43, .25, 12, WHITE, True, PP_ALIGN.CENTER)
        text_box(slide, heading, 1.0, y, width - .72, .35, 17, NAVY, True)
        text_box(slide, body, 1.0, y + .37, width - .72, .52, 13, MUTED)
    if visual:
        slide.shapes.add_picture(str(visual), Inches(8.4), Inches(1.65), width=Inches(4.25))
    return slide


def main():
    visuals = {name: make_visual(f'{name}.png', name) for name in ('dashboard', 'install', 'workflow', 'security')}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(8.5), Inches(-.8), Inches(5.6), Inches(5.6))
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
    text_box(slide, 'Education Center', .8, 1.25, 7.2, .55, 25, RGBColor(112, 233, 219), True, PP_ALIGN.LEFT)
    text_box(slide, 'د MIS سیستم\nد نصب او کارولو لارښود', .8, 1.95, 7.0, 1.35, 33, WHITE, True, PP_ALIGN.LEFT)
    text_box(slide, 'د مرکز مدیر، حسابدار او کارکوونکو لپاره', .8, 3.55, 7.0, .35, 16, RGBColor(210, 231, 237), False, PP_ALIGN.LEFT)
    slide.shapes.add_picture(str(visuals['dashboard']), Inches(7.8), Inches(1.05), width=Inches(4.65))

    bullet_slide(prs, '۱. سیستم څه کار کوي؟', 'دا سیستم د تعلیمي مرکز ټول مهم معلومات په یو ځای کې ساتي.', [
        ('زده‌کوونکي او داخلې', 'زده‌کوونکی ثبت، په صنف کې داخله، د اړیکې معلومات او د فعال حالت اداره.'),
        ('فیس او تادیات', 'بل جوړول، رسید ورکول، پاتې او پوره تادیه شوي فیسونه معلومول.'),
        ('حاضري او پایلې', 'ورځنۍ حاضري، غیابت، امتحان او د زده‌کوونکو پایلې ثبتول.'),
        ('MIS راپورونه', 'د میاشتې مالي راپور، د صنف پاتې فیس او بشپړه حاضري لیدل.'),
    ], visuals['dashboard'])
    bullet_slide(prs, '۲. له نصب مخکې چمتووالی', 'له نصب مخکې دا موارد چک کړئ.', [
        ('کمپیوټر او انټرنېټ', 'Windows 10/11، د انټرنېټ اړیکه یوازې د لومړي نصب او تازه کولو لپاره.'),
        ('Python', 'Python 3.12 یا نوې نسخه نصب کړئ او د Command Prompt له لارې یې وازمویئ: python --version'),
        ('د فایل ځای', 'د Education Center فولډر په خوندي ډرایو کې وساتئ؛ د Desktop پر ځای D: ښه انتخاب دی.'),
        ('احتیاطي کاپي', 'هره ورځ د db.sqlite3 او media فولډر کاپي په USB یا Cloud کې واخلئ.'),
    ], visuals['install'])
    bullet_slide(prs, '۳. د لومړي ځل نصب', 'دا قوماندې په PowerShell کې، د پروژې په فولډر کې اجرا کړئ.', [
        ('۱ — فولډر ته تلل', 'cd "C:\\Path\\To\\Education Center"'),
        ('۲ — اړتیاوې نصب', 'python -m pip install django'),
        ('۳ — ډیټابیس جوړول', 'python manage.py migrate'),
        ('۴ — مدیر جوړول', 'python manage.py createsuperuser  — د کارن نوم او پټنوم خوندي وساتئ.'),
    ], visuals['install'])
    bullet_slide(prs, '۴. سیستم څنګه پرانیزو؟', 'هره ورځ د کار تر پیل مخکې همدا دوه ګامونه وکړئ.', [
        ('سرور چالان کړئ', 'په پروژې فولډر کې: python manage.py runserver'),
        ('براوزر پرانیزئ', 'دا پته ولیکئ: http://127.0.0.1:8000/'),
        ('ننوتل', 'د Login له لارې د مدیر یا کارکوونکي حساب سره داخل شئ.'),
        ('کار پای', 'د ورځې په پای کې له حسابه Logout وکړئ او د ډیټابیس Backup واخلئ.'),
    ], visuals['dashboard'])
    bullet_slide(prs, '۵. د پیل مهم تنظیمات', 'له دې ترتیب څخه مه اوړئ؛ معلومات به منظم پاتې شي.', [
        ('۱ — کارکوونکي', 'مدیر، حسابدار، استقبال او استاد حسابونه او رولونه جوړ کړئ.'),
        ('۲ — کورسونه', 'د کورس نوم، فیس، موده او کچه ثبت کړئ.'),
        ('۳ — صنفونه', 'استاد، پیل نېټه، پای نېټه، خونه او تدریسي ورځې ثبت کړئ.'),
        ('۴ — زده‌کوونکي', 'لومړی زده‌کوونکی ثبت کړئ، بیا یې په صحیح صنف کې داخله کړئ.'),
    ], visuals['workflow'])
    bullet_slide(prs, '۶. د هرې ورځې کاري جریان', 'دا ترتیب د معلوماتو د تېروتنې مخه نیسي.', [
        ('زده‌کوونکی → صنف', 'لومړی Student، بیا Enrollment؛ بې داخله زده‌کوونکي ته بل مه جوړوئ.'),
        ('صنف → بل', 'د بل جوړولو پر وخت هم زده‌کوونکی او هم صحیح صنف انتخاب کړئ.'),
        ('بل → تادیه', 'هره تادیه د رسید نمبر سره ثبت کړئ؛ د بل له اندازې زیاته پیسه نه منل کېږي.'),
        ('حاضري → راپور', 'هر تدریسي ورځ حاضر/غایب ثبت کړئ، بیا د MIS راپور وګورئ.'),
    ], visuals['workflow'])
    bullet_slide(prs, '۷. د حاضرۍ ثبتولو اصول', 'دقیق حاضري د راپور د سمون اساس دی.', [
        ('هره تدریسي ورځ ثبت', 'د استاد له ټولګي وروسته حاضر او غایب زده‌کوونکي سمدستي ثبت کړئ.'),
        ('نه دی ثبت شوی', 'په MIS راپور کې دا حالت مانا لري چې د تدریسي ورځې حاضري لا نه ده ثبت شوې.'),
        ('تدریسي ورځې', 'په Class کې attendance_weekdays له حقیقي مهالویش سره برابر کړئ؛ بېلګه شنبه–چهارشنبه: 5,6,0,1,2.'),
        ('میاشتنی کنټرول', 'د هرې میاشتې په پای کې بشپړ د حاضرۍ لست او CSV فایل واخلئ.'),
    ], visuals['dashboard'])
    bullet_slide(prs, '۸. فیس او MIS راپور تشریح', 'د راپور له لارې په څو ثانیو کې مالي وضعیت وګورئ.', [
        ('پاتې فیسونه', 'هغه بلونه چې لا پوره نه دي تادیه شوي؛ تادیه او پاتې اندازه دواړه ښکاري.'),
        ('پوره تادیه شوي', 'هغه زده‌کوونکي چې د ټاکل شوې میاشتې بل یې بشپړ ورکړی.'),
        ('د صنف راپور', 'د هر صنف پاتې فیس او د کورس د ختمېدو پاتې ورځې.'),
        ('فلټر او صادرول', 'میاشت، صنف او زده‌کوونکی وټاکئ؛ حاضري CSV/Excel ته ښکته کړئ.'),
    ], visuals['dashboard'])
    bullet_slide(prs, '۹. رولونه او د معلوماتو ساتنه', 'هر کارکوونکي ته یوازې اړینه اجازه ورکړئ.', [
        ('مدیر', 'ټول سیستم، راپورونه، کاروونکي او تنظیمات اداره کوي.'),
        ('حسابدار', 'بلونه، تادیات او MIS مالي راپورونه اداره کوي.'),
        ('استقبال او استاد', 'یوازې اړوند زده‌کوونکي، صنفونه او حاضرۍ ته لاسرسی ولري.'),
        ('پټنوم او Backup', 'قوي پټنوم وکاروئ، له بل چا سره یې مه شریکوئ، او هره ورځ Backup واخلئ.'),
    ], visuals['security'])
    bullet_slide(prs, '۱۰. عامې ستونزې او حل', 'د ستونزې پر مهال دا چک‌لېست وکاروئ.', [
        ('سیستم نه پرانیستل کېږي', 'وګورئ runserver چالان دی او پته 127.0.0.1:8000 ده.'),
        ('د صنف پای نېټه', 'که نه ښکاري، د کورس موده او د صنف د پیل نېټه وګورئ.'),
        ('فیس په غلط صنف کې', 'بل مه حذف کوئ؛ د مدیر له لارې یې پر صحیح صنف واړوئ.'),
        ('حاضري نه دی ثبت شوی', 'د صنف تدریسي ورځې او د هماغې ورځې Attendance ثبت وګورئ.'),
    ], visuals['security'])
    bullet_slide(prs, '۱۱. د سپارلو وروستی چک‌لېست', 'تر سپارلو مخکې هره خانه تائید کړئ.', [
        ('مدیر حساب', 'د مدیر کارن نوم او خوندي پټنوم مسئول کس ته وسپارئ.'),
        ('لومړني معلومات', 'کورسونه، استادان، صنفونه، زده‌کوونکي او فیسونه سم ثبت شوي وي.'),
        ('ازمایښتي کار', 'یو زده‌کوونکی، یو بل، یوه تادیه او یوه حاضري د ازموینې لپاره ثبت کړئ.'),
        ('Backup او ملاتړ', 'د Backup ځای وټاکئ او د ستونزې پر مهال د تماس مسئول معلوم کړئ.'),
    ], visuals['workflow'])
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    text_box(slide, 'مننه', .8, 1.8, 11.7, .7, 36, WHITE, True, PP_ALIGN.CENTER)
    text_box(slide, 'Education Center MIS\nمنظم معلومات، روښانه راپورونه، ښه اداره', .8, 2.75, 11.7, 1.0, 20, RGBColor(170, 225, 219), False, PP_ALIGN.CENTER)
    slide.shapes.add_picture(str(visuals['workflow']), Inches(4.2), Inches(4.2), width=Inches(4.9))
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
